"""
Excel Generator service for creating customer reports (.xlsx)
"""
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)

MONEY_FMT = '#,##0'

_TITLE_FILL = PatternFill('solid', fgColor='1F4788')
_HEADER_FILL = PatternFill('solid', fgColor='1F4788')
_ORDER_FILL = PatternFill('solid', fgColor='E3F2FD')    # light blue for orders
_PAYMENT_FILL = PatternFill('solid', fgColor='E8F5E9')  # light green for payments
_TOTALS_FILL = PatternFill('solid', fgColor='DCE6F1')

_THIN = Side(style='thin', color='B0B0B0')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COL_ALIGN = {
    1: Alignment(horizontal='center', vertical='center'),  # Sana
    2: Alignment(horizontal='left', vertical='center'),    # Turi
    3: Alignment(horizontal='center', vertical='center'),  # Mashina
}


class ExcelGenerator:
    """Generate styled .xlsx reports for customers"""

    HEADERS = [
        'Sana', 'Turi', 'Mashina', 'Miqdori (kg)', "Narxi (so'm/kg)",
        "Yo'l Harajati", 'Jami', 'Olingan', 'Qarz',
    ]
    MONEY_COLS = (5, 6, 7, 8, 9)  # E..I

    def __init__(self, temp_dir: Path):
        self.temp_dir = temp_dir
        self.temp_dir.mkdir(exist_ok=True)

    def generate_customer_report(
        self,
        customer_data: Dict,
        combined_data: List[Dict] = None,
        summary: Dict = None
    ) -> Path:
        """
        Generate customer report .xlsx with combined orders and payments

        Args:
            customer_data: Customer information
            combined_data: Combined orders and payments list with cumulative debt
            summary: Summary information (optional)

        Returns:
            Path to generated .xlsx file
        """
        if combined_data is None:
            combined_data = []
        if summary is None:
            summary = {}

        filename = f"report_{customer_data['phone']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = self.temp_dir / filename

        wb = Workbook()
        ws = wb.active
        ws.title = 'Hisobot'
        ws.sheet_view.showGridLines = False

        for col, width in zip('ABCDEFGHI', (12, 24, 13, 13, 15, 14, 15, 15, 15)):
            ws.column_dimensions[col].width = width

        # Title
        ws.merge_cells('A1:I1')
        title = ws['A1']
        title.value = "BUYURTMALAR VA QARZYLIK HISOBOTI"
        title.font = Font(name='Arial', bold=True, size=14, color='FFFFFF')
        title.fill = _TITLE_FILL
        title.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30

        # Customer info
        info_rows = [
            ('Mijoz nomi:', customer_data['name']),
            ('Telefon:', customer_data['phone']),
            ('Manzil:', customer_data['address']),
            ('Vaqt:', datetime.now().strftime('%Y-%m-%d %H:%M')),
        ]
        for i, (label, value) in enumerate(info_rows, start=3):
            ws.cell(row=i, column=1, value=label).font = Font(name='Arial', bold=True, size=11)
            ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
            ws.cell(row=i, column=2, value=value).font = Font(name='Arial', size=11)

        # Starting and current debt
        header_row = 10
        first_data_row = header_row + 1

        if combined_data:
            first = combined_data[0]
            if first['type'] == 'order':
                eski_formula = f'=I{first_data_row}-H{first_data_row}'
            else:
                eski_formula = f'=I{first_data_row}+H{first_data_row}'
            eski_qarzdorlik = eski_formula
            hozirgi_qarz = f'=I{first_data_row + len(combined_data) - 1}'
        else:
            eski_qarzdorlik = summary.get('total_debt', 0)
            hozirgi_qarz = summary.get('total_debt', 0)

        debt_font = Font(name='Arial', bold=True, size=12)
        ws.cell(row=7, column=1, value='Eski qarzdorlik:').font = debt_font
        eski_cell = ws.cell(row=7, column=2, value=eski_qarzdorlik)
        eski_cell.font = debt_font
        eski_cell.number_format = '#,##0 "so\'m"'

        ws.cell(row=8, column=1, value='Hozirgi qarz:').font = debt_font
        hozirgi_cell = ws.cell(row=8, column=2, value=hozirgi_qarz)
        hozirgi_cell.font = debt_font
        hozirgi_cell.number_format = '#,##0 "so\'m"'

        # Transactions table
        if combined_data:
            for col, header in enumerate(self.HEADERS, start=1):
                cell = ws.cell(row=header_row, column=col, value=header)
                cell.font = Font(name='Arial', bold=True, size=10, color='FFFFFF')
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = _BORDER
            ws.row_dimensions[header_row].height = 30

            for idx, transaction in enumerate(combined_data):
                row = first_data_row + idx
                if transaction['type'] == 'order':
                    values = [
                        transaction['date'],
                        transaction['cement_type'],
                        transaction['car_number'] or '-',
                        transaction['quantity'],
                        transaction['price_per_kg'],
                        transaction['road_cost'],
                        transaction['total_sum'],
                        transaction['paid_amount'],
                        transaction['cumulative_debt'],
                    ]
                    fill = _ORDER_FILL
                else:
                    values = [
                        transaction['date'],
                        transaction['payment_type'],
                        '-', '-', '-', '-', '-',
                        transaction['paid_amount'],
                        transaction['cumulative_debt'],
                    ]
                    fill = _PAYMENT_FILL

                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.font = Font(name='Arial', size=10)
                    cell.border = _BORDER
                    cell.fill = fill
                    cell.alignment = _COL_ALIGN.get(
                        col, Alignment(horizontal='right', vertical='center')
                    )
                    if col in self.MONEY_COLS and isinstance(value, (int, float)):
                        cell.number_format = MONEY_FMT

            # Totals row
            last_data_row = first_data_row + len(combined_data) - 1
            totals_row = last_data_row + 1
            total_labels = ['JAMI', '', '', f'=SUM(D{first_data_row}:D{last_data_row})',
                            '', '', f'=SUM(G{first_data_row}:G{last_data_row})',
                            f'=SUM(H{first_data_row}:H{last_data_row})',
                            f'=I{last_data_row}']
            for col, value in enumerate(total_labels, start=1):
                cell = ws.cell(row=totals_row, column=col, value=value)
                cell.font = Font(name='Arial', bold=True, size=10)
                cell.fill = _TOTALS_FILL
                cell.border = _BORDER
                cell.alignment = _COL_ALIGN.get(
                    col, Alignment(horizontal='right', vertical='center')
                )
                if col in self.MONEY_COLS:
                    cell.number_format = MONEY_FMT
        else:
            ws.cell(row=header_row, column=1, value='Buyurtmalar topilmadi').font = \
                Font(name='Arial', italic=True, size=11)

        ws.freeze_panes = f'A{first_data_row}'
        wb.save(filepath)
        logger.info(f"Generated Excel report: {filepath}")
        return filepath
