"""Self-check for ExcelGenerator: builds report with sample data, verifies formulas recalc."""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from bot.services.excel_generator import ExcelGenerator  # noqa: E402

customer = {
    'name': 'Test Mijoz LLC',
    'phone': '998934701803',
    'address': 'Toshkent sh., Test ko\'cha 1',
}

combined = [
    {'type': 'order', 'date': '2026-01-10', 'cement_type': 'M400', 'car_number': '01A123BC',
     'quantity': 5000, 'price_per_kg': 1500, 'road_cost': 100000, 'total_sum': 7400000,
     'paid_amount': 4000000, 'remaining_debt': 3400000, 'cumulative_debt': 3400000},
    {'type': 'order', 'date': '2026-02-15', 'cement_type': 'M500', 'car_number': '01B456DE',
     'quantity': 3000, 'price_per_kg': 1600, 'road_cost': 50000, 'total_sum': 4750000,
     'paid_amount': 2000000, 'remaining_debt': 2750000, 'cumulative_debt': 6150000},
    {'type': 'payment', 'date': '2026-03-01', 'payment_type': "Pul ko'chirish",
     'paid_amount': 3000000, 'cumulative_debt': 3150000},
    {'type': 'order', 'date': '2026-04-20', 'cement_type': 'M400', 'car_number': '-',
     'quantity': 2000, 'price_per_kg': 1550, 'road_cost': 0, 'total_sum': 3100000,
     'paid_amount': 1000000, 'remaining_debt': 2100000, 'cumulative_debt': 5250000},
]

tmp = Path('/tmp/xlsx_test')
tmp.mkdir(exist_ok=True)
gen = ExcelGenerator(tmp)
path = gen.generate_customer_report(customer, combined, {'total_debt': 5250000})
print(f'Created: {path}')

# no-data case (different phone so filename unique)
customer2 = dict(customer, phone='998900000000')
path2 = gen.generate_customer_report(customer2, [], {'total_debt': 0})
print(f'Created: {path2}')

# verify formulas evaluate via LibreOffice recalc
import subprocess
script = Path.home() / '.claude/plugins/cache/anthropic-agent-skills/document-skills/41bbe19d1a1a/skills/xlsx/scripts/recalc.py'
r = subprocess.run([sys.executable, str(script), str(path), '60'], capture_output=True, text=True)
print(r.stdout)
assert '"status": "success"' in r.stdout, 'recalc failed'

from openpyxl import load_workbook
wb = load_workbook(path, data_only=True)
ws = wb.active
# eski qarzdorlik = first order: cumulative(3400000) - paid(4000000) = -600000
assert ws['B7'].value == -600000, f'eski qarzdorlik wrong: {ws["B7"].value}'
# hozirgi qarz = last cumulative = 5250000
assert ws['B8'].value == 5250000, f'hozirgi qarz wrong: {ws["B8"].value}'
# totals row: header at row 10, first data row 11
tr = 10 + 1 + len(combined)
assert ws.cell(row=tr, column=4).value == 10000, f'qty total wrong: {ws.cell(row=tr, column=4).value}'
# jami sum = 7400000+4750000+3100000 = 15250000
assert ws.cell(row=tr, column=7).value == 15250000, f'total_sum wrong: {ws.cell(row=tr, column=7).value}'
# olingan total = 4000000+2000000+3000000+1000000 = 10000000
assert ws.cell(row=tr, column=8).value == 10000000, f'paid total wrong'
# qarz total = last cumulative 5250000
assert ws.cell(row=tr, column=9).value == 5250000, f'debt total wrong'
print('ALL CHECKS PASSED')
