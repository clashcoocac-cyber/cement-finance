"""
Excel export views: styled .xlsx reports generated with openpyxl.
Each view mirrors the corresponding page's filtered data.
"""
from datetime import date

from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Sum
from django.http import HttpResponse
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finance.filters import CustomerFilter, OrderFilter, PaymentFilter
from finance.models import CementType, Customer, Order, PaymentHistory

MONEY_FMT = "#,##0"

_TITLE_FILL = PatternFill("solid", fgColor="1F4788")
_HEADER_FILL = PatternFill("solid", fgColor="1F4788")
_TOTALS_FILL = PatternFill("solid", fgColor="DCE6F1")
_ORDER_FILL = PatternFill("solid", fgColor="E3F2FD")
_PAYMENT_FILL = PatternFill("solid", fgColor="E8F5E9")

_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_LEFT = Alignment(horizontal="left", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")
_CENTER = Alignment(horizontal="center", vertical="center")


def _title_row(ws, title, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    c.fill = _TITLE_FILL
    c.alignment = _CENTER
    ws.row_dimensions[1].height = 28


def _header_row(ws, headers, row):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
    ws.row_dimensions[row].height = 26


def _style_cell(cell, bold=False, fill=None, money=False, align=None, color=None):
    cell.font = Font(name="Arial", size=10, bold=bold, color=color)
    if fill:
        cell.fill = fill
    cell.border = _BORDER
    cell.alignment = align or _RIGHT
    if money:
        cell.number_format = MONEY_FMT


def _money(v):
    return v if v is not None else 0


def _response(wb, name):
    from datetime import datetime

    from django.http import HttpResponse

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="{name}_{ts}.xlsx"'
    wb.save(resp)
    return resp


class OrderExportView(LoginRequiredMixin, View):
    """Buyurtmalar sahifasi eksporti (order.html jadvali)."""

    def get(self, request):
        query_params = request.GET.copy()
        if not query_params.get("customer_id"):
            query_params["date_from"] = query_params.get("date_from") or date.today()
            query_params["date_to"] = query_params.get("date_to") or date.today()

        filtered_orders = OrderFilter(
            query_params,
            queryset=Order.objects.order_by("order_date").select_related(
                "customer", "cement_type"
            ),
        ).qs

        customer = None
        if request.GET.get("customer_id"):
            try:
                customer = Customer.objects.get(id=request.GET.get("customer_id"))
            except Customer.DoesNotExist:
                pass

        if customer:
            # Mirror the page: combine orders and payments like OrderView
            from finance.views import OrderView

            date_from = query_params.get("date_from")
            date_to = query_params.get("date_to")
            payments = OrderView._get_payment_data(OrderView(), customer.id, date_from, date_to)
            rows = OrderView._prepare_combined_data(
                OrderView(), filtered_orders, payments, customer
            )
            rows = OrderView._calculate_cumulative_debt(OrderView(), rows, customer)
        else:
            rows = filtered_orders

        wb = Workbook()
        ws = wb.active
        ws.title = "Buyurtmalar"
        ws.sheet_view.showGridLines = False
        for col, width in zip("ABCDEFGHIJ", (12, 22, 14, 12, 14, 14, 14, 15, 15, 15)):
            ws.column_dimensions[col].width = width

        _title_row(ws, "BUYURTMALAR HISOBOTI — {}".format(
            customer.name if customer else "Barcha mijozlar"), 10)

        if customer:
            info = f"Mijoz: {customer.name}   Eski qarz: {customer.total_debt:,} so'm"
        else:
            info = f"Sana: {date.today().strftime('%Y-%m-%d')}"
        ws.cell(row=2, column=1, value=info).font = Font(
            name="Arial", italic=True, size=10
        )

        headers = [
            "Sana", "Mijoz", "Turi", "Mashina", "Miqdori (kg)",
            "Narxi (so'm/kg)", "Yo'l Xarajati", "Jami", "Olingan", "Qarz",
        ]
        _header_row(ws, headers, 4)

        row = 5
        for order in rows:
            is_payment = isinstance(order, dict) and order.get("type") == "payment"
            if is_payment:
                values = [
                    order["order_date"], str(order["customer"]),
                    "To'lov", order.get("comment") or "", None, None, None,
                    None, _money(order["paid_amount"]), _money(order["remaining_debt"]),
                ]
                fill = _PAYMENT_FILL
            elif isinstance(order, dict):
                values = [
                    order["order_date"], str(order["customer"]),
                    str(order["cement_type"]) if order["cement_type"] else "",
                    order["car_number"], _money(order["quantity"]),
                    _money(order["price_per_kg"]), _money(order["road_cost"]),
                    _money(order["total_sum"]), _money(order["paid_amount"]),
                    _money(order["remaining_debt"]),
                ]
                fill = _ORDER_FILL
            else:
                values = [
                    order.order_date,
                    order.customer.name if order.customer else "",
                    order.cement_type.name if order.cement_type else "",
                    order.car_number, _money(order.quantity),
                    _money(order.price_per_kg), _money(order.road_cost),
                    _money(order.total_sum), _money(order.paid_amount),
                    _money(order.remaining_debt),
                ]
                fill = _ORDER_FILL

            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                money = col >= 5 and v is not None
                _style_cell(
                    c, fill=fill,
                    align=_LEFT if col in (2, 3, 4) else _CENTER if col == 1 else _RIGHT,
                    money=money,
                    color="D32F2F" if col == 10 and v and v > 0 else
                          "388E3C" if col == 10 else None,
                )
            row += 1

        # Totals
        if isinstance(rows[0], dict) if rows else False:
            total_quantity = sum(o["quantity"] for o in rows if o.get("type") == "order")
            total_paid = sum(o["paid_amount"] for o in rows)
            total_debt = customer.total_debt
        else:
            total_quantity = sum(o.quantity for o in rows)
            total_paid = sum(o.paid_amount for o in rows)
            total_debt = sum(o.remaining_debt for o in rows)

        totals = ["JAMI:", "", "", "", total_quantity,
                  "", "", "", total_paid, total_debt]
        for col, v in enumerate(totals, start=1):
            c = ws.cell(row=row, column=col, value=v)
            _style_cell(c, bold=True, fill=_TOTALS_FILL,
                        align=_LEFT if col == 1 else _RIGHT, money=col >= 5)
        row += 2
        ws.cell(row=row, column=1, value="Hozirgi qarz:").font = Font(
            name="Arial", bold=True, size=11
        )
        c = ws.cell(row=row, column=2, value=total_debt)
        c.font = Font(name="Arial", bold=True, size=11)
        c.number_format = MONEY_FMT

        ws.freeze_panes = "A5"
        return _response(wb, "buyurtmalar_hisoboti")


class CustomerExportView(LoginRequiredMixin, View):
    """Mijozlar sahifasi eksporti (customer.html jadvali)."""

    def get(self, request):
        customers = CustomerFilter(
            request.GET, Customer.objects.order_by("-total_debt")
        ).qs
        total_debt = customers.aggregate(total=Sum("total_debt"))["total"] or 0

        wb = Workbook()
        ws = wb.active
        ws.title = "Mijozlar"
        ws.sheet_view.showGridLines = False
        for col, width in zip("ABCDE", (6, 28, 18, 35, 18)):
            ws.column_dimensions[col].width = width

        _title_row(ws, "MIJOZLAR RO'YXATI", 5)
        _header_row(ws, ["#", "Nomi", "Telefon", "Manzil", "Qolgan Qarz"], 3)

        row = 4
        for i, customer in enumerate(customers, start=1):
            values = [i, customer.name, customer.phone, customer.address,
                      _money(customer.total_debt)]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                _style_cell(
                    c, money=(col == 5),
                    align=_LEFT if col in (2, 3, 4) else _CENTER if col == 1 else _RIGHT,
                    color="D32F2F" if col == 5 and customer.total_debt > 0
                          else "388E3C" if col == 5 else None,
                )
            row += 1

        for col, v in enumerate(["JAMI:", "", "", "", total_debt], start=1):
            c = ws.cell(row=row, column=col, value=v)
            _style_cell(c, bold=True, fill=_TOTALS_FILL,
                        align=_LEFT if col == 1 else _RIGHT, money=(col == 5))

        ws.freeze_panes = "A4"
        return _response(wb, "mijozlar_royxati")


class DebtExportView(LoginRequiredMixin, View):
    """Qarzlar sahifasi eksporti (debt.html jadvali)."""

    def get(self, request):
        selected_year = request.GET.get("year", str(date.today().year))
        try:
            selected_year = int(selected_year)
        except (ValueError, TypeError):
            selected_year = date.today().year

        date_from = request.GET.get("date_from", date.today().strftime("%Y-%m-%d"))
        date_to = request.GET.get("date_to", date.today().strftime("%Y-%m-%d"))

        query_params = request.GET.copy()
        query_params["date_from"] = date_from
        query_params["date_to"] = date_to

        payments = PaymentFilter(
            query_params,
            PaymentHistory.objects.filter(paid_at__year=selected_year).order_by(
                "-paid_at"
            ),
        ).qs

        wb = Workbook()
        ws = wb.active
        ws.title = "Qarz to'lovlari"
        ws.sheet_view.showGridLines = False
        for col, width in zip("ABCDE", (16, 24, 18, 16, 35)):
            ws.column_dimensions[col].width = width

        _title_row(ws, f"QARZ TO'LOVLARI — {selected_year}-yil", 5)
        _header_row(
            ws,
            ["Sana / Vaqt", "Mijoz", "To'langan Summa", "To'lov turi", "Izoh"],
            3,
        )

        row = 4
        for payment in payments:
            values = [
                payment.paid_at.strftime("%Y-%m-%d %H:%M"),
                payment.customer.name if payment.customer else "",
                _money(payment.amount),
                payment.get_payment_type_display() or "-",
                payment.comment or "",
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                _style_cell(
                    c, money=(col == 3),
                    align=_LEFT if col in (2, 4, 5) else _CENTER if col == 1 else _RIGHT,
                    color="388E3C" if col == 3 else None,
                )
            row += 1

        total_amount = sum(p.amount for p in payments)
        for col, v in enumerate(["JAMI:", "", total_amount, "", ""], start=1):
            c = ws.cell(row=row, column=col, value=v)
            _style_cell(c, bold=True, fill=_TOTALS_FILL,
                        align=_LEFT if col == 1 else _RIGHT, money=(col == 3))

        ws.freeze_panes = "A4"
        return _response(wb, "qarz_tovlovlari")


class CementTypeExportView(LoginRequiredMixin, View):
    """Sement turlari sahifasi eksporti (cement_type.html jadvali)."""

    def parse_month(self, month_str):
        try:
            year, month_num = month_str.split("-")
            return int(year), int(month_num)
        except ValueError:
            today = date.today()
            return today.year, today.month

    def get(self, request):
        selected_month = request.GET.get("month") or date.today().strftime("%Y-%m")
        year, month = self.parse_month(selected_month)

        from django.db import models

        month_sum = models.Sum(
            models.Case(
                models.When(
                    orders__order_date__year=year,
                    orders__order_date__month=month,
                    then="orders__quantity",
                ),
                default=0,
            )
        )
        cement_types = CementType.objects.annotate(total_quantity=month_sum)

        wb = Workbook()
        ws = wb.active
        ws.title = "Sement turlari"
        ws.sheet_view.showGridLines = False
        for col, width in zip("ABC", (24, 16, 26)):
            ws.column_dimensions[col].width = width

        _title_row(ws, f"SEMENT TURLARI — {selected_month}", 3)
        _header_row(
            ws,
            ["Turi", "Rang", "Jami sotilgan miqdori (kg)"],
            3,
        )

        row = 4
        for ct in cement_types:
            ws.cell(row=row, column=1, value=ct.name)
            _style_cell(ws.cell(row=row, column=1), align=_LEFT)
            ws.cell(row=row, column=2, value=ct.get_color_display())
            _style_cell(ws.cell(row=row, column=2), align=_CENTER)
            c = ws.cell(row=row, column=3, value=_money(ct.total_quantity))
            _style_cell(c, money=True)
            row += 1

        total_quantity = sum(ct.total_quantity or 0 for ct in cement_types)
        for col, v in enumerate(["Jami:", "", total_quantity], start=1):
            c = ws.cell(row=row, column=col, value=v)
            _style_cell(c, bold=True, fill=_TOTALS_FILL,
                        align=_LEFT if col == 1 else _RIGHT, money=(col == 3))

        ws.freeze_panes = "A4"
        return _response(wb, "sement_turlari")
